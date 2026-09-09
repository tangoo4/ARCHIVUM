"""Creación de copias de respaldo independientes de una temporada."""

from pathlib import Path
import shutil

from openpyxl import load_workbook

from core.nomenclatura import obtener_nombre_base


FILA_INICIO_DATOS = 2
COL_TOMO = "A"
INTERVALO_RESPALDO = 5
CARPETA_RESPALDO = "respaldo"


def _ruta_respaldo(ruta_excel, nombre):
    ruta_excel = Path(ruta_excel)
    carpeta = ruta_excel.parent / CARPETA_RESPALDO
    nombre_base = obtener_nombre_base(ruta_excel)
    return carpeta / f"{nombre_base} {nombre}{ruta_excel.suffix}"


def _copiar_si_no_existe(ruta_excel, ruta_destino):
    ruta_excel = Path(ruta_excel)
    if not ruta_excel.exists():
        raise FileNotFoundError(f"No existe el archivo principal: {ruta_excel}")

    if ruta_destino.exists():
        return None

    ruta_destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ruta_excel, ruta_destino)
    return ruta_destino


def crear_respaldo_periodico(ruta_excel, numero_tomo):
    """Crea el respaldo T005, T010, ... tras guardar un nuevo tomo.

    Devuelve la ruta creada o None cuando no corresponde crear respaldo
    o cuando ese punto de respaldo ya existe.
    """
    numero_tomo = int(numero_tomo)
    if numero_tomo <= 0 or numero_tomo % INTERVALO_RESPALDO != 0:
        return None

    destino = _ruta_respaldo(ruta_excel, f"T{numero_tomo:03d}")
    return _copiar_si_no_existe(ruta_excel, destino)


def obtener_ultimo_tomo(ruta_excel):
    """Devuelve el número del último tomo real guardado en el Excel."""
    ruta_excel = Path(ruta_excel)
    wb = load_workbook(ruta_excel, read_only=True, data_only=False)
    ws = wb.active
    ultimo_tomo = None

    try:
        for fila in range(FILA_INICIO_DATOS, ws.max_row + 1):
            try:
                tomo = int(ws[f"{COL_TOMO}{fila}"].value)
            except (TypeError, ValueError):
                continue
            ultimo_tomo = tomo
    finally:
        wb.close()

    return ultimo_tomo


def crear_respaldo_final(ruta_excel):
    """Crea una copia FINAL del Excel ya cerrado."""
    ultimo_tomo = obtener_ultimo_tomo(ruta_excel)
    if ultimo_tomo is None:
        raise ValueError("Todavía no existen tomos medidos.")

    destino = _ruta_respaldo(ruta_excel, f"FINAL T{ultimo_tomo:03d}")
    return _copiar_si_no_existe(ruta_excel, destino)
