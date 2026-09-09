"""Lectura común de los tomos válidos de una temporada."""

from pathlib import Path
from openpyxl import load_workbook

FILA_INICIO_DATOS = 2


def es_tomo_real(valor):
    try:
        int(valor)
        return True
    except (TypeError, ValueError):
        return False


def leer_tomos_temporada(ruta_excel, data_only=False):
    wb = load_workbook(Path(ruta_excel), data_only=data_only)
    ws = wb.active
    tomos = []
    for fila in range(FILA_INICIO_DATOS, ws.max_row + 1):
        if not es_tomo_real(ws[f"A{fila}"].value):
            continue
        tomos.append({
            "fila_excel": fila,
            "tomo": ws[f"A{fila}"].value,
            "anio": ws[f"B{fila}"].value,
            "matriz_inicio": ws[f"C{fila}"].value,
            "fecha_inicio": ws[f"D{fila}"].value,
            "matriz_final": ws[f"E{fila}"].value,
            "fecha_final": ws[f"F{fila}"].value,
            "medida": ws[f"G{fila}"].value,
            "observaciones": ws[f"H{fila}"].value or "",
        })
    return tomos
