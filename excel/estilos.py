"""Estilos Excel compartidos de ARCHIVUM."""

import copy
from openpyxl.styles import Alignment, PatternFill

COLOR_VERDE_FONDO = "C6EFCE"
COLOR_VERDE_TEXTO = "006100"
COLOR_AMARILLO_FONDO = "FFEB9C"
COLOR_AMARILLO_TEXTO = "9C5700"
COLOR_ROJO_FONDO = "FFC7CE"
COLOR_ROJO_TEXTO = "9C0006"

FILL_VERDE = PatternFill("solid", fgColor=COLOR_VERDE_FONDO)
FILL_AMARILLO = PatternFill("solid", fgColor=COLOR_AMARILLO_FONDO)
FILL_ROJO = PatternFill("solid", fgColor=COLOR_ROJO_FONDO)
ALINEACION_CENTRADA = Alignment(horizontal="center", vertical="center")


def estilo_estado_medida(medida):
    if medida == "?":
        return FILL_AMARILLO, COLOR_AMARILLO_TEXTO
    return FILL_VERDE, COLOR_VERDE_TEXTO


def aplicar_estado_fila(ws, fila, medida, columnas="ABCDEFGH"):
    fill, color_texto = estilo_estado_medida(medida)
    for columna in columnas:
        celda = ws[f"{columna}{fila}"]
        celda.fill = copy.copy(fill)
        fuente = copy.copy(celda.font)
        fuente.color = color_texto
        celda.font = fuente
        celda.alignment = copy.copy(ALINEACION_CENTRADA)
